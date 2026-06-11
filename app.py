import streamlit as st
import google.generativeai as genai
from openai import OpenAI
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import time

# ==================== 1. 初始化与页面配置 ====================
st.set_page_config(page_title="AI 智能学习助手", layout="wide")

# 初始化所有需要持久化的状态
if "analyzed" not in st.session_state:
    st.session_state.analyzed = False
if "outline" not in st.session_state:
    st.session_state.outline = ""
if "raw_text" not in st.session_state:
    st.session_state.raw_text = ""
if "quiz_data" not in st.session_state:
    st.session_state.quiz_data = None
if "quiz_type" not in st.session_state:
    st.session_state.quiz_type = ""
if "show_answers" not in st.session_state:
    st.session_state.show_answers = False
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

# ==================== 2. 文件解析工具类 ====================
def extract_text(uploaded_file):
    file_ext = uploaded_file.name.split(".")[-1].lower()
    text = ""
    if file_ext == "pdf":
        reader = PdfReader(uploaded_file)
        for page in reader.pages:
            t = page.extract_text()
            if t: text += t + "\n"
    elif file_ext in ["docx", "doc"]:
        doc = Document(uploaded_file)
        for para in doc.paragraphs: text += para.text + "\n"
    elif file_ext in ["pptx", "ppt"]:
        prs = Presentation(uploaded_file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text: text += shape.text + "\n"
    elif file_ext == "txt":
        text = uploaded_file.read().decode("utf-8", errors="ignore")
    return text

# ==================== 3. AI 调用路由 ====================
def call_ai(provider, api_key, prompt):
    if provider == "DeepSeek (国内直连)":
        client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com")
        response = client.chat.completions.create(
            model="deepseek-chat",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3
        )
        return response.choices[0].message.content
    else:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel('gemini-1.5-flash')
        response = model.generate_content(prompt)
        return response.text

# ==================== 4. 侧边栏布局 (与手绘图一致) ====================
with st.sidebar:
    st.header("🔑 1. API 接口配置")
    ai_provider = st.selectbox("选择模型引擎", ["DeepSeek (国内直连)", "Google Gemini"])
    api_key = st.text_input(f"输入 {ai_provider} Key", type="password")
    
    st.header("📁 2. 文件导入区")
    if not st.session_state.analyzed:
        uploaded_file = st.file_uploader("支持 PDF/Word/PPT", type=["pdf", "docx", "pptx", "txt"])
        start_btn = st.button("🚀 点击开始分析", use_container_width=True)
    else:
        st.success("✅ 文件已锁定分析")
        if st.button("🔄 重新上传文件", use_container_width=True):
            for key in st.session_state.keys(): del st.session_state[key]
            st.rerun()

# 后台解析逻辑
if not st.session_state.analyzed and 'start_btn' in locals() and start_btn:
    if not api_key or not uploaded_file:
        st.error("请确保 API Key 和文件都已准备就绪！")
    else:
        with st.spinner("正在解析文件内容并生成大纲..."):
            text = extract_text(uploaded_file)
            st.session_state.raw_text = text
            prompt = f"请对以下文本进行结构化拆解，输出大纲（包含：引言、理论背景、理论核心重点、案例分析、对策建议、结论）：\n\n{text}"
            st.session_state.outline = call_ai(ai_provider, api_key, prompt)
            st.session_state.analyzed = True
            st.rerun()

# ==================== 5. 右侧主界面布局 (按照手绘图重构) ====================
# 页面顶部标题
st.markdown(f"<h1 style='text-align: center;'>🎓 智能全格式课件与论文大纲重构 Agent</h1>", unsafe_allow_html=True)

if st.session_state.analyzed:
    # A/B/C 三个选项卡切换
    tab_a, tab_b, tab_c = st.tabs(["📋 A 文件大纲", "✍️ B 习题练习", "💬 C 问题讨论"])

    # --- Tab A: 文件大纲 ---
    with tab_a:
        st.markdown("### 📄 课件/论文结构化拆解")
        st.markdown(st.session_state.outline)

    # --- Tab B: 习题练习 ---
    with tab_b:
        st.markdown("### ✏️ 智能习题测试")
        # 习题类型选择
        q_type = st.radio("选择题目形式：", ["选择题 (单选10道)", "填空题 (10道)", "论述大题 (3道)"], horizontal=True)
        
        if st.button(f"生成{q_type}"):
            with st.spinner("正在根据文件内容出题..."):
                count = "10" if "10" in q_type else "3"
                q_prompt = f"""
                基于以下文本内容，生成{q_type}。
                要求：
                1. 题目必须严谨且贴合原文。
                2. 必须在所有题目生成完毕后，最后再给出答案。
                3. 使用 [ANSWERS] 字符串作为题目和答案的分隔符。
                
                原文：{st.session_state.raw_text[:4000]} 
                """
                raw_quiz = call_ai(ai_provider, api_key, q_prompt)
                if "[ANSWERS]" in raw_quiz:
                    parts = raw_quiz.split("[ANSWERS]")
                    st.session_state.quiz_data = {"q": parts[0], "a": parts[1]}
                else:
                    st.session_state.quiz_data = {"q": raw_quiz, "a": "AI未按格式返回答案，请在讨论区咨询。"}
                st.session_state.show_answers = False

        if st.session_state.quiz_data:
            st.divider()
            st.markdown(st.session_state.quiz_data["q"])
            
            if not st.session_state.show_answers:
                if st.button("✅ 已完成所有题目，查看答案"):
                    st.session_state.show_answers = True
                    st.rerun()
            else:
                st.markdown("---")
                st.markdown("#### 🔑 参考答案")
                st.success(st.session_state.quiz_data["a"])
                if st.button("🙈 隐藏答案"):
                    st.session_state.show_answers = False
                    st.rerun()

    # --- Tab C: 问题讨论 ---
    with tab_c:
        st.markdown("### 💬 深度知识研讨")
        # 显示对话历史
        for msg in st.session_state.chat_history:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])
        
        # 用户输入
        if chat_input := st.chat_input("向 AI 提问或讨论习题..."):
            with st.chat_message("user"):
                st.markdown(chat_input)
            st.session_state.chat_history.append({"role": "user", "content": chat_input})
            
            with st.chat_message("assistant"):
                with st.spinner("AI 正在思考..."):
                    context = f"你是一个学术助教。基于文件：{st.session_state.raw_text[:2000]}。回答提问：{chat_input}"
                    reply = call_ai(ai_provider, api_key, context)
                    st.markdown(reply)
            st.session_state.chat_history.append({"role": "assistant", "content": reply})

else:
    # 未上传文件时的欢迎界面
    st.info("👈 请在左侧侧边栏配置 API 并上传你的课件或论文文件开始分析。")
    st.image("https://img.icons8.com/illustrations/ul/480/searching.png", width=400) # 装饰性插画
